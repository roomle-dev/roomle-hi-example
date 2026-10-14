from datetime import date
from hashlib import sha256
from pathlib import Path
import json
import re
import zipfile
import xml.etree.ElementTree as ElementTree

import pypandoc
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = ROOT / ".github" / "skills" / "kitchen-planning"
OUTPUT = ROOT / "Kitchen-Planning-Skill-Documentation.docx"
SOURCES = [
    ("SKILL.md", "Full Skill Definition", "skill-definition"),
    ("references/planning-reference.md", "Full Planning Reference", "planning-reference"),
    ("references/source-catalog.md", "Full Source Catalog", "source-catalog"),
]
OVERVIEW = """# Documentation Guide

## Purpose and audience

This document explains how to use the kitchen-planning skill and includes the
complete contents of its three Markdown files, formatted for reading in Word.
It is intended for kitchen planners, reviewers, and people integrating the skill
into an AI-assisted planning workflow.

The skill supports residential kitchen planning with Germany/EU defaults,
millimetres for dimensions, and EUR for costs. It is vendor-neutral and must be
adapted to the actual jurisdiction, site, products, and household.

## Contents

- [Using the skill](#using-the-skill)
- [Full skill definition, including metadata](#skill-definition)
- [Full planning reference](#planning-reference)
- [Full source catalog](#source-catalog)

The headings also support Word's Navigation Pane. Links between the included
files lead to the corresponding chapter; external source links remain clickable.

## Document status

This is a documentation snapshot, not an executable skill, a product catalog,
or a professional installation approval. The Markdown files in the workspace
remain the maintained originals. The chapters reproduce their full content;
Markdown syntax is rendered as headings, lists, tables, code, and links. The
skill's YAML metadata is included explicitly as a code block.

The source catalog is included here for completeness. Its presence in this
document does not change the skill's optional-loading policy.

# Using the Skill

## Installation and invocation

Keep the following structure in the workspace:

```text
.github/skills/kitchen-planning/
    SKILL.md
    references/
        planning-reference.md
        source-catalog.md
```

In a compatible VS Code Copilot environment, invoke `/kitchen-planning` or ask
for a kitchen-planning task that matches the skill's description. Skill discovery
depends on the host's support and configuration. In another AI application, load
the skill through that application's instruction/resource mechanism; a Word file
alone does not install or activate the skill.

## Which files to load

| File | Role | Loading policy |
| --- | --- | --- |
| SKILL.md | Main workflow, intake, constraints, cabinet assignment, task trials, and deliverables. | Load when the skill is invoked. |
| references/planning-reference.md | Dimensions with qualifications, 40 check categories, and worked examples. | Read before applying numerical guidance or reviewing installation risks; use its checklist as directed by the skill. |
| references/source-catalog.md | Ten source entries, URLs, publication details, and research notes. | Optional: load for source verification, provenance/citation questions, or guidance updates, not routine planning. |

Local regulations, exact product instructions, and professional checks may still
be required for a project. Making the bibliography optional does not make those
checks optional.

## Example requests

**New layout:** "Use /kitchen-planning to develop alternatives for my kitchen in
Germany. I will provide a measured floor plan, fixed service points, appliance
models, household needs, and budget. Identify missing information before claiming
the layout is dimensionally checked."

**Storage review:** "Review this layout for dishwasher unloading. Assign dining
cutlery, plates, bowls, cups, and cooking utensils to cabinets. Test the dishwasher
and receiving drawers open together, and explain the routes and trade-offs."

**Evidence review:** "Explain the source and limitations of the suggested
preparation width between sink and hob. Consult the source catalog and distinguish
ergonomic guidance from installation requirements."

## Planning sequence

1. Gather the household brief, inventory, measured geometry, services, and products.
2. Separate hard constraints from recommendations and preferences; preserve unknowns.
3. Arrange food storage, preparation, cooking, cleaning, and clean-storage zones.
4. Assign actual contents before selecting cabinet dimensions and drawer fittings.
5. Compare feasible alternatives using open-state task trials and walkable routes.
6. Deliver the layout, storage schedule, constraint report, and professional-review actions.

Plan for daily activities rather than cabinet symmetry alone. A useful sink-to-hob
preparation surface is an ergonomic goal, not a universal EU legal distance.
Nearby dish storage is useful only when it remains reachable with the dishwasher
and receiving storage open. Cooking tools generally belong at their point of use.

## Expected results and limits

The output should explain what goes in each cabinet, how everyday tasks work,
which facts are verified, and which constraints remain unresolved. Readiness
progresses through `concept`, `dimensionally_checked`, and
`ready_for_professional_order_review`. None of these replaces required professional
installation approval or separate authorization to place an order.

The skill is written guidance, not a deterministic geometry validator. It does
not provide an appliance database, certify legal compliance, or establish hidden
services from photographs. Use authoritative geometry/catalog tools where
available, and label estimates or missing information honestly.

## Maintenance and rebuilding

Edit the maintained Markdown files, then regenerate this document with:

```text
python scripts/build-kitchen-documentation.py
```

The builder requires `python-docx`, `pypandoc`, and an available Pandoc executable
(for example, provided by `pypandoc_binary`). It reads the latest three files,
generates the Word document, and checks that their text content is represented.
It does not modify the skill files. Review the generated document before external
publication, especially after introducing new Markdown constructs or wide tables.
"""


def parse_markdown(text):
    return json.loads(pypandoc.convert_text(text, "json", format="gfm"))


def visit(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from visit(child)
    elif isinstance(value, list):
        for child in value:
            yield from visit(child)


def inline_text(value):
    if isinstance(value, list):
        return "".join(inline_text(child) for child in value)
    if not isinstance(value, dict):
        return ""
    kind = value.get("t")
    content = value.get("c")
    if kind == "Str":
        return content
    if kind in ("Space", "SoftBreak", "LineBreak"):
        return " "
    if kind == "Code":
        return content[1]
    if kind in ("Link", "Image"):
        return inline_text(content[1])
    return inline_text(content)


def normalize(text):
    return re.sub(r"\s+", " ", text).strip()


def style_document():
    document = Document(OUTPUT)
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(18)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name in ("Body Text", "First Paragraph", "Compact"):
        if style_name in document.styles:
            document.styles[style_name].font.name = "Calibri"
            document.styles[style_name].font.size = Pt(10.5)
    for level in range(1, 5):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor.from_string("176B65")
        style.font.size = Pt({1: 21, 2: 16, 3: 12, 4: 11}[level])
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
    document.styles["Heading 1"].paragraph_format.page_break_before = True
    for style_name in ("Source Code", "Verbatim Char"):
        if style_name in document.styles:
            document.styles[style_name].font.name = "Consolas"
            document.styles[style_name].font.size = Pt(8)
    for table in document.tables:
        table.autofit = False
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "BCCDCB")
            borders.append(border)
        table._tbl.tblPr.append(borders)
        column_count = len(table.columns)
        proportions = [0.25, 0.75] if column_count == 2 else [0.23, 0.32, 0.45]
        if table.cell(0, 0).text == "ID":
            proportions = [0.13, 0.23, 0.64]
        if len(proportions) != column_count:
            proportions = [1 / column_count] * column_count
        for column, proportion in zip(table.columns, proportions):
            column.width = Mm(174 * proportion)
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for row_index, row in enumerate(table.rows):
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell, proportion in zip(row.cells, proportions):
                cell.width = Mm(174 * proportion)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.keep_with_next = row_index == 0
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        if row_index == 0:
                            run.bold = True
                if row_index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "E5F1EF")
                    cell._tc.get_or_add_tcPr().append(shading)
    section.header.paragraphs[0].text = "KITCHEN PLANNING | Skill Documentation"
    footer = section.footer.paragraphs[0]
    footer.add_run("Documentation snapshot | Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(OUTPUT)


def main():
    source_bytes = {relative: (SKILL_ROOT / relative).read_bytes() for relative, _, _ in SOURCES}
    expected_files = {relative for relative, _, _ in SOURCES}
    actual_files = {path.relative_to(SKILL_ROOT).as_posix() for path in SKILL_ROOT.rglob("*") if path.is_file()}
    if actual_files != expected_files:
        raise ValueError("Skill files changed; update the documentation source list before building.")
    combined = parse_markdown(OVERVIEW)
    combined["meta"] = {
        "title": {"t": "MetaString", "c": "Kitchen Planning Skill"},
        "subtitle": {"t": "MetaString", "c": "Usage Guide and Complete Skill Files"},
        "date": {"t": "MetaString", "c": date.today().isoformat()},
        "lang": {"t": "MetaString", "c": "en-GB"},
    }
    source_documents = []
    target_anchors = {(SKILL_ROOT / relative).resolve(): anchor for relative, _, anchor in SOURCES}
    for relative, title, anchor in SOURCES:
        text = source_bytes[relative].decode("utf-8-sig")
        if relative == "SKILL.md":
            frontmatter = re.match(r"\A---\r?\n.*?\r?\n---(?:\r?\n|$)", text, re.DOTALL)
            if not frontmatter:
                raise ValueError("Missing skill metadata.")
            text = "```yaml\n" + frontmatter[0].rstrip() + "\n```\n\n" + text[frontmatter.end():]
        parsed = parse_markdown(text)
        source_documents.append(parsed)
        chapter = parse_markdown(f"# {title}\n\nSource file: `{relative}`\n")
        chapter["blocks"][0]["c"][1][0] = anchor
        combined["blocks"].extend(chapter["blocks"])
        for node in visit(parsed["blocks"]):
            if node.get("t") == "Header":
                node["c"][0] = min(node["c"][0] + 1, 6)
                node["c"][1][0] = anchor + "-" + node["c"][1][0]
            elif node.get("t") == "Link":
                target = node["c"][2][0]
                if target.startswith("./"):
                    resolved = ((SKILL_ROOT / relative).parent / target).resolve()
                    node["c"][2][0] = "#" + target_anchors[resolved]
        combined["blocks"].extend(parsed["blocks"])
    pypandoc.convert_text(json.dumps(combined), "docx", format="json", outputfile=str(OUTPUT))
    style_document()
    with zipfile.ZipFile(OUTPUT) as package:
        if package.testzip() is not None:
            raise ValueError("Corrupt Word package.")
        for member in package.namelist():
            if member.endswith((".xml", ".rels")):
                ElementTree.fromstring(package.read(member))
        document_xml = ElementTree.fromstring(package.read("word/document.xml"))
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    paragraphs = [
        "".join(
            (element.text or "") if element.tag == qn("w:t") else " "
            for element in paragraph.iter()
            if element.tag in (qn("w:t"), qn("w:br"), qn("w:tab"))
        )
        for paragraph in document_xml.findall(".//w:p", namespaces)
    ]
    all_text = normalize(" ".join(paragraphs))
    checked_blocks = 0
    for parsed in source_documents:
        for node in visit(parsed["blocks"]):
            kind = node.get("t")
            if kind in ("Para", "Plain"):
                expected = inline_text(node["c"])
            elif kind == "Header":
                expected = inline_text(node["c"][2])
            elif kind == "CodeBlock":
                expected = node["c"][1]
            else:
                continue
            if normalize(expected) not in all_text:
                raise ValueError(f"Missing source content: {expected[:120]}")
            checked_blocks += 1
    bookmarks = {node.get(qn("w:name")) for node in document_xml.findall(".//w:bookmarkStart", namespaces)}
    for hyperlink in document_xml.findall(".//w:hyperlink", namespaces):
        anchor = hyperlink.get(qn("w:anchor"))
        if anchor and anchor not in bookmarks:
            raise ValueError(f"Unresolved internal document link: {anchor}")
    for relative, original in source_bytes.items():
        if sha256((SKILL_ROOT / relative).read_bytes()).digest() != sha256(original).digest():
            raise ValueError(f"Source changed during generation: {relative}")
    reopened = Document(OUTPUT)
    print(f"Created: {OUTPUT.name} ({OUTPUT.stat().st_size:,} bytes)")
    print(f"PASS: {len(SOURCES)} complete files, {checked_blocks} source text blocks, "
          f"{len(reopened.tables)} tables, internal links, XML package, and unchanged source files.")


if __name__ == "__main__":
    main()
